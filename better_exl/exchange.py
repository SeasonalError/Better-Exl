"""Lossless input snapshots and inspectable analysis exports."""

import csv, io, json, re, zipfile
from pathlib import Path
from .templates import blank_sheet, col, uid
from .engine import analyze, MAX_ROWS, MAX_COLUMNS
from .expressions import CONSTANTS, FUNCTIONS


def safe_cell(v):
    t = "" if v is None else str(v)
    if t.lstrip().startswith(("=", "+", "@")) or (
        t.lstrip().startswith("-") and not re.fullmatch(r"\s*-[\d.eE+\-]+\s*", t)
    ):
        return "'" + t
    return t


def csv_text(rows):
    out = io.StringIO(newline="")
    w = csv.writer(out)
    for r in rows:
        w.writerow([safe_cell(x) for x in r])
    return out.getvalue()


def import_table(
    blob, filename, header=0, delimiter="auto", decimal_comma=False, tab=""
):
    header = int(header)
    if not 0 <= header <= 100:
        raise ValueError("Header row must be 0–100.")
    names = []
    if filename.lower().endswith(".xlsx"):
        import openpyxl

        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            if sum(x.file_size for x in z.infolist()) > 60_000_000:
                raise ValueError("Expanded workbook exceeds 60 MB.")
        wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=False)
        names = wb.sheetnames
        ws = wb[tab] if tab in names else wb[names[0]]
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS + header:
                raise ValueError("Too many rows for this version.")
            rows.append(["" if v is None else str(v) for v in r])
        tab = ws.title
        wb.close()
    else:
        try:
            text = blob.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Resave this file as UTF-8 CSV/TSV/TXT.")
        if delimiter == "auto":
            try:
                delimiter = (
                    csv.Sniffer()
                    .sniff(
                        "\n".join(text.splitlines()[header : header + 20]),
                        delimiters=",\t;|",
                    )
                    .delimiter
                )
            except csv.Error:
                delimiter = "\t" if filename.endswith((".tsv", ".txt")) else ","
        if delimiter not in (",", "\t", ";", "|"):
            raise ValueError("Unsupported separator.")
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if len(rows) <= header or not rows[header]:
        raise ValueError("No header at that row.")
    headers = rows[header]
    body = [r for r in rows[header + 1 :] if any(str(v).strip() for v in r)]
    if not 1 <= len(headers) <= MAX_COLUMNS or len(body) > MAX_ROWS:
        raise ValueError(f"Import limit: {MAX_ROWS:,} rows and {MAX_COLUMNS} columns.")
    columns = []
    used = set(CONSTANTS) | set(FUNCTIONS)
    for i, title in enumerate(headers):
        m = re.fullmatch(r"(.*?)\s*\[([^\]]*)\]\s*", str(title))
        name, un = (
            (m[1].strip(), m[2].strip())
            if m
            else (str(title).strip() or f"Column {i+1}", "")
        )
        base = re.sub(r"[^A-Za-z0-9_]", "_", name).strip("_")[:30] or f"col{i+1}"
        if not base[0].isalpha():
            base = "v_" + base
        key = base
        j = 2
        while key in used:
            key = f"{base}_{j}"
            j += 1
        used.add(key)
        columns.append(col(key, name, un))
    s = blank_sheet(Path(filename).stem, columns)
    for r in body:
        if len(r) > len(columns):
            raise ValueError(
                "A data row is wider than the header. Check delimiter/decimal settings."
            )
        s["rows"].append(
            dict(
                id=uid(),
                values={
                    c["key"]: (
                        str(r[i]).replace(",", ".") if decimal_comma else str(r[i])
                    )
                    for i, c in enumerate(columns)
                    if i < len(r)
                },
                uncertainties={},
                included=True,
                note="",
            )
        )
    s["source"] = f"Imported from {filename}" + (f" · {tab}" if tab else "")
    s["importSnapshot"] = {
        "headers": headers,
        "rows": body,
        "filename": filename,
        "sheet": tab,
    }
    return {"sheet": s, "workbookSheets": names, "selectedSheet": tab}


def data_csv(s, a=None, processed=True):
    a = a or analyze(s)
    cols = (
        s["columns"] if processed else [c for c in s["columns"] if not c.get("formula")]
    )
    headers = ["row", "included", "note"]
    for c in cols:
        un = a["units"].get(c["key"], c.get("unit", ""))
        headers += [f"{c['name']} [{un}]", f"u({c['name']}) [{un}]"]
    rows = [headers]
    for i, r in enumerate(s["rows"]):
        row = [i + 1, r.get("included", True), r.get("note", "")]
        for c in cols:
            k = c["key"]
            row += [
                a["values"][k][i] if processed else r.get("values", {}).get(k, ""),
                (
                    a["uncertainties"][k][i]
                    if processed
                    else r.get("uncertainties", {}).get(k, "")
                ),
            ]
        rows.append(row)
    return csv_text(rows)


def report(p, s, a):
    e = lambda x: str(x).replace("|", "\\|").replace("\n", " ")
    lines = [
        f"# {p['name']}",
        f"\nDataset: {s['name']}",
        f"\nSource: {s.get('source','')}",
        f"\nObjective: {p.get('objective','')}",
        f"\nApparatus: {p.get('apparatus','')}",
        "\n## Methods and assumptions",
        s.get("notes", ""),
        "\nStandard uncertainties are one sigma. Blank uncertainty is unknown. Propagation uses first-order J Σ Jᵀ with central finite differences, including shared source dependencies.",
        "\n## Columns",
        "| Symbol | Name | Unit | Formula | Uncertainty |",
        "|---|---|---|---|---|",
    ]
    for c in s["columns"]:
        lines.append(
            "| "
            + " | ".join(
                e(x)
                for x in [
                    c["key"],
                    c["name"],
                    a["units"].get(c["key"], c.get("unit", "")),
                    c.get("formula") or "Measured",
                    f"{c.get('uncertainty','')} ({c.get('uncertaintyMode','absolute')})",
                ]
            )
            + " |"
        )
    if a["fit"]:
        f = a["fit"]
        lines += [
            "\n## Fit",
            f"`{f['expression']}` · {f['method']} · n={f['n']} · degrees of freedom={f['dof']}",
            "| Parameter | Value | Standard uncertainty | Unit |",
            "|---|---:|---:|---|",
        ]
        for pa in f["params"]:
            lines.append(
                f"| {e(pa['name'])} | {pa['value']} | {pa['stderr'] if pa['stderr'] is not None else 'Fixed'} | {e(pa.get('unit',''))} |"
            )
        lines += [
            f"\nRMSE={f['rmse']}; descriptive R²={f['r2']}; χ²={f['chi2']}; reduced χ²={f['reducedChi2']}.",
            "Intervals are local covariance approximations; R² is not a test of physical validity. Fit bands describe mean-curve uncertainty, not observation prediction intervals.",
        ]
    if a.get("fitError"):
        lines += ["\nFit failed: " + a["fitError"]]
    if a["theory"]:
        lines += [
            "\n## Independent theory",
            f"Equation: `{s['theory']['expression']}`. Entered constants: {json.dumps(s['theory']['params'],ensure_ascii=False)}.",
            f"RMSE={a['theory']['rmse']}. Theory-parameter uncertainty is not included in the overlay.",
        ]
    if a.get("theoryError"):
        lines += ["\nTheory failed: " + a["theoryError"]]
    lines += (
        ["\n## Diagnostics"]
        + ["- " + x for x in a["diagnostics"]]
        + ["\n## Observations", p.get("notes", ""), "\n## Analysis history"]
        + [f"- {x.get('time','')}: {x.get('action','')}" for x in p.get("log", [])]
        + [
            "\nGenerated by Better Exl 1.0.0. Inspect project.json and analysis.json for complete inputs and covariance. Round results to appropriate significant figures before reporting."
        ]
    )
    return "\n".join(lines)


def python_script(s):
    return (
        '"""Run in the Better-Exl environment from the repository root."""\nimport json, sys\nfrom pathlib import Path\nsys.path.insert(0, str(Path.cwd()))\nfrom better_exl.engine import analyze\nfrom better_exl.exchange import data_csv\nsheet=json.loads('
        + repr(json.dumps(s, ensure_ascii=False))
        + ')\nresult=analyze(sheet)\nPath("reproduced-analysis.json").write_text(json.dumps(result,indent=2),encoding="utf-8")\nPath("reproduced-data.csv").write_text(data_csv(sheet,result),encoding="utf-8")\nprint(json.dumps(result.get("fit"),indent=2))\n'
    )


def bundle(p, s):
    a = analyze(s)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in {
            "project.json": json.dumps(p, indent=2, ensure_ascii=False),
            "raw.csv": data_csv(s, a, False),
            "processed.csv": data_csv(s, a),
            "analysis.json": json.dumps(a, indent=2, ensure_ascii=False),
            "report.md": report(p, s, a),
            "reproduce.py": python_script(s),
            "README.txt": "Raw CSV contains entered uncertainty overrides only. Defaults/modes are in project.json; processed.csv contains evaluated standard uncertainties. Run reproduce.py in the installed Better Exl environment. Export figures separately.\n",
        }.items():
            z.writestr(name, text)
    buf.seek(0)
    return buf
